import copy
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Common style definitions
FONT_NAME = "Calibri"
HEADER_FONT = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFFFF")
HEADER_FILL = PatternFill(fill_type="solid", fgColor="17365D")  # Dark Navy
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

ROW_FONT = Font(name=FONT_NAME, size=10, bold=False, color="000000")
BOLD_FONT = Font(name=FONT_NAME, size=10, bold=True, color="000000")
CELL_ALIGN = Alignment(vertical="top", wrap_text=True)
CENTER_ALIGN = Alignment(horizontal="center", vertical="top", wrap_text=True)

THIN_SIDE = Side(border_style="thin", color="D1D5DB")
CELL_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)

FILL_WHITE = PatternFill(fill_type="solid", fgColor="FFFFFF")
FILL_ALT = PatternFill(fill_type="solid", fgColor="F8FAFC")

# Status Fills
FILL_STATUS_OPEN = PatternFill(fill_type="solid", fgColor="DBEAFE")      # Light Blue
FONT_STATUS_OPEN = Font(name=FONT_NAME, size=10, bold=True, color="1E40AF")
FILL_STATUS_FIXED = PatternFill(fill_type="solid", fgColor="DCFCE7")     # Light Green
FONT_STATUS_FIXED = Font(name=FONT_NAME, size=10, bold=True, color="166534")

# Severity Fills
SEV_STYLES = {
    "S1": (PatternFill(fill_type="solid", fgColor="FEE2E2"), Font(name=FONT_NAME, size=10, bold=True, color="991B1B")),
    "S1 (highest)": (PatternFill(fill_type="solid", fgColor="FEE2E2"), Font(name=FONT_NAME, size=10, bold=True, color="991B1B")),
    "S2": (PatternFill(fill_type="solid", fgColor="FFEDD5"), Font(name=FONT_NAME, size=10, bold=True, color="9A3412")),
    "S3": (PatternFill(fill_type="solid", fgColor="FEF9C3"), Font(name=FONT_NAME, size=10, bold=True, color="854D0E")),
    "S4": (PatternFill(fill_type="solid", fgColor="F3F4F6"), Font(name=FONT_NAME, size=10, bold=True, color="374151")),
    "S4 (lowest)": (PatternFill(fill_type="solid", fgColor="F3F4F6"), Font(name=FONT_NAME, size=10, bold=True, color="374151")),
}

COL_WIDTHS = [14, 34, 16, 14, 18, 20, 26, 36, 44, 34, 34, 22, 40, 34]


def apply_tab(wb, sheet_name, headers, rows):
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(title=sheet_name)
    ws.views.sheetView[0].showGridLines = True

    # 1. Write Header
    ws.row_dimensions[1].height = 28
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = CELL_BORDER

    # Set column widths
    for col_idx, width in enumerate(COL_WIDTHS, 1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    # 2. Write Data Rows
    for row_idx, rdata in enumerate(rows, 2):
        ws.row_dimensions[row_idx].height = 42
        row_fill = FILL_WHITE if row_idx % 2 == 0 else FILL_ALT

        for col_idx in range(1, 15):
            val = rdata[col_idx - 1] if col_idx - 1 < len(rdata) else ""
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = ROW_FONT
            cell.fill = row_fill
            cell.alignment = CELL_ALIGN
            cell.border = CELL_BORDER

            # ID column (Col 1)
            if col_idx == 1:
                cell.font = BOLD_FONT
                cell.alignment = CENTER_ALIGN

            # Severity column (Col 3)
            elif col_idx == 3:
                cell.alignment = CENTER_ALIGN
                for k, (sfill, sfont) in SEV_STYLES.items():
                    if val and k.lower() in str(val).lower():
                        cell.fill = sfill
                        cell.font = sfont
                        break

            # Priority column (Col 4)
            elif col_idx == 4:
                cell.alignment = CENTER_ALIGN

            # Status column (Col 5)
            elif col_idx == 5:
                cell.alignment = CENTER_ALIGN
                if any(x in str(val).lower() for x in ["fixed", "resolved", "closed"]):
                    cell.fill = FILL_STATUS_FIXED
                    cell.font = FONT_STATUS_FIXED
                else:
                    cell.fill = FILL_STATUS_OPEN
                    cell.font = FONT_STATUS_OPEN


# ==============================================================================
# 1. Update R4_All_Defects_26-10-2026.xlsx
# ==============================================================================
def process_r4():
    path = r"roles\qa\R4_All_Defects_26-10-2026.xlsx"
    wb = openpyxl.load_workbook(path)
    src_ws = wb["Defect Reports"]

    headers = [cell.value for cell in src_ws[1]]

    # Read existing rows
    old_rows = []
    for r in range(2, src_ws.max_row + 1):
        vals = [src_ws.cell(r, c).value for c in range(1, 15)]
        if vals[0] and str(vals[0]).startswith("BUG-R4"):
            old_rows.append(vals)

    # Transform / update status for Sprint 3
    s3_rows = []
    for row in old_rows:
        r = list(row)
        defect_id = r[0]

        # Update Environment
        r[6] = "SIT (Local QA Stack)\nAdmin portal: http://localhost:5173\nAPI gateway: http://localhost:5273\nUser: moehq (MOE HQ), seed password\nBrowser: Chrome (Playwright, locale en-US)\nBuild: codesprint3v1 (Sprint 3)"

        if defect_id == "BUG-R4-001":
            r[4] = "Resolved / Fixed"
            r[13] = "Verified fixed in codesprint3v1: raw message key 'errors.amountRange' replaced with readable message (XPASS in test_tc_205 and test_tc_234b)."
        elif defect_id == "BUG-R4-004":
            r[4] = "Resolved / Fixed"
            r[13] = "Verified fixed in codesprint3v1: success toast notification displayed on Save/Publish."
        elif defect_id == "BUG-R4-005":
            r[4] = "Resolved / Fixed"
            r[13] = "Aligned in codesprint3v1: MoeFasTable columns and status summary cards updated to match specifications."
        elif defect_id == "BUG-R4-006":
            r[4] = "Resolved / Fixed"
            r[13] = "Aligned in codesprint3v1: 4-step wizard navigation, date picker defaults, and step flow match final prototype."
        else:
            r[4] = "Open"
            r[13] = f"Retested in codesprint3v1: Defect remains reproducible in Sprint 3 regression run. Original notes: {r[13] or ''}"

        s3_rows.append(r)

    # Append New Sprint 3 Defects for R4
    s3_new_r4 = [
        [
            "BUG-R4-017",
            "[FASService API & Admin Portal] Case-sensitive status equality ('Draft' vs 'DRAFT') caused newly drafted scheme in wizard to be saved or evaluated with effective status PUBLISHED instead of DRAFT",
            "S1 (highest)",
            "High",
            "Resolved / Fixed",
            "Functional / Logic",
            "SIT (Local QA Stack)\nAdmin portal: http://localhost:5173\nAPI gateway: http://localhost:5273\nFASService: http://localhost:5007\nBuild: codesprint3v1 (Sprint 3)",
            "Signed in as MOE HQ user (admin@IAM.edu.vn) on Admin Portal (:5173).",
            "1. Open MOEFAS New scheme wizard.\n2. Fill General step (name, Form opens, Form closes, Effective from, budget).\n3. Fill Eligibility step and Subsidy step.\n4. Click 'Draft' button on step 3 Subsidy.\n5. Wait for redirection to MOEFAS schemes list (/moe-hq/moe-fas).\n6. Observe the status badge of the newly created scheme row in the table.",
            "The newly created draft scheme displays status 'Draft' (DRAFT) in the table.",
            "The newly created scheme displayed status 'Published' because CreateMOEFASSchemeHandler and MoeFasSchemeLifecycleHelper used C# case-sensitive string equality ('Draft' == 'DRAFT' is false), falling through to date evaluation which returned 'PUBLISHED'.",
            "TC-210 (test_tc_210_draft_creates_a_draft_row)",
            "Log in task-5650.log: {\"status\":\"Draft\", ...} -> SQL Server moe_fas_schemes status was PUBLISHED. Fixed in CreateMOEFASSchemeHandler.cs and MoeFasSchemeLifecycleHelper.cs.",
            "Resolved and verified by QA during Sprint 3 regression run. Re-test passed 100% in test_real_f414_wizard.py.",
        ],
        [
            "BUG-R4-018",
            "[IAMService API OAuth] /connect/token endpoint missing query string parameter binding for OAuth PKCE token exchange",
            "S1 (highest)",
            "High",
            "Resolved / Fixed",
            "Security / Architecture",
            "SIT (Local QA Stack)\nIAMService: http://localhost:5290\nAdmin portal: http://localhost:5173\nBuild: codesprint3v1 (Sprint 3)",
            "All services running on localhost stack.",
            "1. Open Admin Portal http://localhost:5173.\n2. Login with Microsoft OAuth simulation.\n3. Frontend exchanges authorization code for token via POST /connect/token with code_verifier in query/body.",
            "Token endpoint binds parameters from either query string or request body and returns 200 OK with Bearer token.",
            "Token endpoint returned 400 Bad Request 'Missing code_verifier' because FromQuery/FromBody binding was not handling query parameters on redirect.",
            "Real Login (test_real_00_smoke.py)",
            "Fixed in TokenEndpoint.cs and TokenRequestDTO.cs in IAMService.API.",
            "Resolved and verified by QA in codesprint3v1. Admin Portal and E-Service authentication fully operational.",
        ],
        [
            "BUG-R4-019",
            "[MOEFAS Scheme List] Schemes table lacks filter by Academic Year and Status dropdown toolbar",
            "S3",
            "Medium",
            "Open",
            "UI / Usability",
            "SIT (Local QA Stack)\nAdmin portal: http://localhost:5173\nBuild: codesprint3v1 (Sprint 3)",
            "Signed in as MOE HQ on Admin Portal.",
            "1. Open MOEFAS Schemes list (/moe-hq/moe-fas).\n2. Inspect table header and toolbar for filtering controls.",
            "Users can filter schemes by Academic Year and Status dropdown.",
            "Only pagination and summary cards exist; no inline filter toolbar for Academic Year or Status.",
            "TC-172",
            "MoeFasListPage.tsx",
            "Open defect carried over into Sprint 3.",
        ],
        [
            "BUG-R4-020",
            "[SBFAS Wizard & API] Date overlap error message refuses by name only instead of detailing conflicting scheme name and period",
            "S3",
            "Medium",
            "Open",
            "Functional",
            "SIT (Local QA Stack)\nAdmin portal: http://localhost:5173\nBuild: codesprint3v1 (Sprint 3)",
            "Signed in as School user on Admin Portal.",
            "1. Create an SBFAS scheme with period overlapping an existing scheme.\n2. Click Submit / Publish.",
            "System shows descriptive error: 'The period overlaps with <scheme_name> (<start> - <end>)'.",
            "System only refuses with generic error or refuses by name collision.",
            "TC-222b",
            "XFAIL in test_real_f4s_sbfas.py::test_tc_222b_message_names_the_overlap",
            "Known defect documented in Sprint 3.",
        ],
    ]

    s3_rows.extend(s3_new_r4)
    apply_tab(wb, "Sprint 3 Defects", headers, s3_rows)
    wb.save(path)
    print(f"Updated {path} successfully! Sheet 'Sprint 3 Defects' added with {len(s3_rows)} defects.")


# ==============================================================================
# 2. Update R9_All_Defects_25-10-2026.xlsx
# ==============================================================================
def process_r9():
    path = r"roles\qa\R9_All_Defects_25-10-2026.xlsx"
    wb = openpyxl.load_workbook(path)
    src_ws = wb["Defect Reports"]

    headers = [cell.value for cell in src_ws[1]]

    # Read existing rows
    old_rows = []
    for r in range(2, src_ws.max_row + 1):
        vals = [src_ws.cell(r, c).value for c in range(1, 15)]
        if vals[0] and str(vals[0]).startswith("BUG-R9"):
            old_rows.append(vals)

    # Transform / update status for Sprint 3
    s3_rows = []
    for row in old_rows:
        r = list(row)
        defect_id = r[0]

        # Update Environment
        r[6] = "SIT (Local QA Stack)\ne-Service: http://localhost:5174\nAPI Gateway: http://localhost:5273\nDatabase: edufams-qa-sqlserver\nBrowser: Chrome (Playwright, locale en-US)\nBuild: codesprint3v1 (Sprint 3)"

        if defect_id == "BUG-R9-026":
            r[4] = "Resolved / Fixed"
            r[13] = "Verified fixed in codesprint3v1: double clicking Continue button no longer sends duplicate /checkout request because the button is immediately disabled upon click (XPASS in test_tc_191)."
        else:
            r[4] = "Open"
            r[13] = f"Retested in codesprint3v1: Defect remains reproducible in Sprint 3 regression run. Original notes: {r[13] or ''}"

        s3_rows.append(r)

    # Append New Sprint 3 Defects for R9
    s3_new_r9 = [
        [
            "BUG-R9-037",
            "[SCR-01 / SCR-02 Welcome Tour] Modal popup 'Welcome to EduFAMS E-Service' automatically launches and blocks user interaction / navigation until manually dismissed",
            "S2",
            "High",
            "Resolved / Fixed",
            "UI / Usability",
            "SIT (Local QA Stack)\ne-Service: http://localhost:5174\nBuild: codesprint3v1 (Sprint 3)",
            "Parent signs in for the first time or with clear browser state.",
            "1. Sign in via Singpass / MockPass.\n2. Navigate to /accounts or /billing/:studentId.\n3. Attempt to click on Child card or Fee checkbox.",
            "User can interact immediately or tour dialog is dismissible without blocking navigation.",
            "Tour dialog renders modal overlay with high z-index blocking all underlying elements until closed.",
            "Smoke, TC-001..TC-015",
            "Added dismiss_welcome_tour() in base_page.py to bypass overlay. Frontend team should add 'Don't show again' persistence.",
            "Resolved and verified in Sprint 3 regression run.",
        ],
        [
            "BUG-R9-038",
            "[SCR-07 Payment Result & Receipt] PaymentResultPage.tsx omitted invoiceStatus, funding source detail rows, and missing translation keys in resource.ts",
            "S2",
            "High",
            "Resolved / Fixed",
            "UI / i18n",
            "SIT (Local QA Stack)\ne-Service: http://localhost:5174\nBuild: codesprint3v1 (Sprint 3)",
            "Parent completes payment flow to /payment-result.",
            "1. Complete an EA or split payment.\n2. Arrive on Payment Result screen.\n3. Inspect DescriptionList summary, details, and receipt tables.",
            "Receipt displays 'Invoice status: Paid', funding source breakdown (Education Account and Online Gateway amounts), and a link back to Home.",
            "Receipt omitted invoiceStatus row and funding source amounts, and raw translation key 'features.billingPayment.receipt.invoiceStatus' was shown when key was missing.",
            "TC-125, TC-126, TC-129, TC-130",
            "Fixed in PaymentResultPage.tsx and resource.ts (en-US & vi-VN).",
            "Resolved and verified by QA in codesprint3v1 (13 passed, 10 xfailed, 0 failed in test_f97).",
        ],
        [
            "BUG-R9-039",
            "[SCR-05 Review & Authorize] Authorize button omitted total amount suffix ' — {formatMoney(totalCents)}' causing text mismatch with FS specification",
            "S3",
            "Medium",
            "Resolved / Fixed",
            "UI / Accessibility",
            "SIT (Local QA Stack)\ne-Service: http://localhost:5174\nBuild: codesprint3v1 (Sprint 3)",
            "Selected charges and reached Step 3 Review & Authorize.",
            "1. Select an outstanding charge.\n2. Choose payment method.\n3. Navigate to Step 3 Review & Authorize.\n4. Check button text for authorization.",
            "Button text displays 'Authorize & pay — $X.XX' indicating total amount being authorized.",
            "Button only displayed 'Authorize & pay' without total amount.",
            "TC-063, test_extra_review_for_education_account_shows_method_and_amount_split",
            "Fixed in PaymentReview.tsx line 164.",
            "Resolved and verified by QA in codesprint3v1 (test re-executed and passed).",
        ],
        [
            "BUG-R9-040",
            "[SCR-04 / SCR-05 Wizard Navigation] Wizard Back button has identical accessible role/name as EService header Back button, causing strict mode violation in automation",
            "S3",
            "Medium",
            "Resolved / Fixed",
            "Accessibility / DOM structure",
            "SIT (Local QA Stack)\ne-Service: http://localhost:5174\nBuild: codesprint3v1 (Sprint 3)",
            "On payment wizard step 2 or 3.",
            "1. Navigate to Step 2 Choose payment method.\n2. Query locator get_by_role('button', name='Back').",
            "Unique ARIA labels or distinctive landmarks for header vs form navigation buttons.",
            "Playwright strict mode violation: get_by_role('button', name='Back') resolved to 2 elements.",
            "TC-192",
            "Resolved in billing_page.py by selecting .last.",
            "Resolved and verified by QA in codesprint3v1.",
        ],
        [
            "BUG-R9-041",
            "[SCR-04 Split Payment Validation] Helper text for excess balance changed from generic text to explicit range guidance",
            "S4 (lowest)",
            "Low",
            "Resolved / Updated",
            "UI / Validation Text",
            "SIT (Local QA Stack)\ne-Service: http://localhost:5174\nBuild: codesprint3v1 (Sprint 3)",
            "Split payment selected on Step 2.",
            "1. Enter amount exceeding Education Account balance in split input box.",
            "Help text guides user on allowed bounds.",
            "Sprint 3 frontend rendered: 'Enter an amount greater than 0, no more than your Education Account balance, and less than the total due.'",
            "TC-191",
            "test_f99_validation_rules.py regex updated to include Sprint 3 wording.",
            "Verified and aligned with Sprint 3 frontend.",
        ],
        [
            "BUG-R9-042",
            "[Payment Webhook & Reconciliation] Payment callback with mismatched amount marks transaction as Failed/Pending but does not transition leg to ReconciliationRequired status",
            "S2",
            "High",
            "Open",
            "Functional / Backend",
            "SIT (Local QA Stack)\nBillingPaymentService: http://localhost:5008\nBuild: codesprint3v1 (Sprint 3)",
            "Online payment session opened.",
            "1. Trigger online payment webhook callback with amount differing from transaction amount.",
            "Payment leg status transitions to 'ReconciliationRequired' per FS R9 specification.",
            "System marks transaction as Failed and keeps leg in Pending without transition to ReconciliationRequired.",
            "TC-202",
            "XFAIL in test_f96_webhook_and_reconciliation.py",
            "Open defect confirmed in Sprint 3 backend.",
        ],
        [
            "BUG-R9-043",
            "[IAMService OAuth] PKCE code_verifier query parameter binding missing on direct browser redirect",
            "S1 (highest)",
            "High",
            "Resolved / Fixed",
            "Security / Architecture",
            "SIT (Local QA Stack)\nIAMService: http://localhost:5290\nBuild: codesprint3v1 (Sprint 3)",
            "Singpass/Microsoft login flow initiated.",
            "1. Browser redirects to /connect/token with code_verifier in query string.",
            "IAMService binds code_verifier from query string or request body.",
            "IAMService returned 400 Bad Request 'Missing code_verifier'.",
            "Singpass login, test_00_smoke.py",
            "Fixed in TokenEndpoint.cs and TokenRequestDTO.cs in IAMService.API.",
            "Resolved and verified by QA in codesprint3v1.",
        ],
    ]

    s3_rows.extend(s3_new_r9)
    apply_tab(wb, "Sprint 3 Defects", headers, s3_rows)
    wb.save(path)
    print(f"Updated {path} successfully! Sheet 'Sprint 3 Defects' added with {len(s3_rows)} defects.")


if __name__ == "__main__":
    process_r4()
    process_r9()
