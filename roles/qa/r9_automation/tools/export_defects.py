"""Xuất toàn bộ 65 case Failed (Known Defect) của R9 sang sheet Defect Reports,
sheet Evidence (kèm 23 ảnh PNG nhúng trực tiếp) và điền mã BUG-R9-### vào cột Q (Defect ID)
của sheet 'Test cases ' theo đúng định dạng mẫu trong roles/qa."""

from pathlib import Path
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

EV_DIR = Path(r"D:\AIQA\Agileproject\roles\qa\r9_automation\artifacts\evidence")

ENV_TEXT = (
    "SIT (Local QA Stack)\n"
    "e-Service: http://localhost:5174\n"
    "API Gateway: http://localhost:5273\n"
    "Database: edufams-qa-sqlserver (SQL Server 2022)\n"
    "Browser: Chrome (latest), Windows 11\n"
    "Build: QA_BE / develop (22 Oct 2026)"
)

DEFECTS = [
    {
        "id": "BUG-R9-001",
        "summary": "[SCR-01 Choose child] Child cards do not show school, level, class or outstanding invoice label",
        "severity": "S3",
        "priority": "Medium",
        "status": "Open",
        "type": "UI",
        "env": ENV_TEXT,
        "pre": "Parent (S7890123F, Tan Ah Kow) is signed in through Singpass and has linked children with and without outstanding charges.",
        "steps": (
            "1. Sign in to the e-Service portal as Parent S7890123F.\n"
            "2. Open the Education Accounts / Choose Child screen (/accounts, SCR-01).\n"
            "3. Inspect each child card for school name, level/cohort, class, and the outstanding/no-outstanding invoices indicator."
        ),
        "expected": (
            "Each child entry on SCR-01 displays the student name, school, level (cohort), class, and a label indicating whether the child has outstanding invoices or no outstanding invoices (TC-004, TC-005)."
        ),
        "actual": (
            "Each card only displays the student name, account number (e.g., 'Account EDU-00502'), status ('Status: Active') and a 'View fees for <Child>' link. School name, level, class and the outstanding/no-outstanding label are not shown."
        ),
        "tcs": ["TC-004", "TC-005"],
        "ev_num": 1,
        "caption": "SCR-01 (/accounts): Thẻ học sinh chỉ hiện tên, mã EDU và trạng thái Active; thiếu trường, khối, lớp và nhãn khoản phí chưa trả",
        "ev_file": "evidence_01_BUG-R9-001.png",
        "notes": "Root cause: AccountsPage in apps/e-service only renders wallet summary fields from EducationAccountService and does not fetch or display school/level/class or billing outstanding summary.",
    },
    {
        "id": "BUG-R9-002",
        "summary": "[API & SCR-02 Access Control] Unlinked student ID returns HTTP 200 empty list instead of 403 Forbidden and UI does not show access denied message",
        "severity": "S1 (highest)",
        "priority": "Critical",
        "status": "Open",
        "type": "Security",
        "env": ENV_TEXT,
        "pre": "Parent (S7890123F) is signed in and requests outstanding charges or opens /billing/{studentId} for a student ID belonging to another family.",
        "steps": (
            "1. Sign in as Parent S7890123F.\n"
            "2. Call GET /api/billing/v1/students/00000000-0000-4000-8000-000000000001/outstanding-charges (TC-012).\n"
            "3. Navigate directly to http://localhost:5174/billing/00000000-0000-4000-8000-000000000001 in the browser (TC-016)."
        ),
        "expected": (
            "1. The API refuses the request with HTTP 403 Forbidden and writes a security log entry (TC-012).\n"
            "2. Navigating to /billing/{unlinkedStudentId} returns 403 and displays 'You do not have access to this student's records.' without showing wallet or billing controls (TC-016)."
        ),
        "actual": (
            "1. GET /api/billing/v1/students/{unlinkedStudentId}/outstanding-charges returns HTTP 200 OK with {\"isSuccess\":true,\"statusCode\":200,\"data\":[]}.\n"
            "2. Navigating to /billing/{unlinkedStudentId} opens the billing screen with 'Outstanding fees SGD 0.00' instead of showing 'You do not have access to this student's records.'"
        ),
        "tcs": ["TC-012", "TC-016"],
        "ev_num": 2,
        "caption": "API trả HTTP 200 [] thay vì 403 Forbidden và UI mở trang /billing/{unlinkedStudentId} thay vì chặn truy cập",
        "ev_file": "evidence_02_BUG-R9-002.png",
        "notes": "Root cause: StudentChargesController and GetOutstandingChargesHandler do not verify guardian-student linkage against IAM/EducationAccount before querying charges.",
    },
    {
        "id": "BUG-R9-003",
        "summary": "[SCR-02 Outstanding fees] Flat charge table is displayed instead of expandable invoices grouped by fee type with GST and earliest due date",
        "severity": "S2",
        "priority": "High",
        "status": "Open",
        "type": "UI",
        "env": ENV_TEXT,
        "pre": "Parent is signed in and opens SCR-02 (/billing/01999001-0000-7000-8000-000000000001) for a child with multiple outstanding fee lines.",
        "steps": (
            "1. Open /billing/01999001-0000-7000-8000-000000000001 (SCR-02).\n"
            "2. Inspect the list of outstanding items for invoice numbers (INV-YYYY-NNNNNN), earliest due date across lines, expandable fee-type groups (Tuition, School, Miscellaneous), and GST breakdown."
        ),
        "expected": (
            "1. Open invoices are listed with their invoice number and the earliest due date of their lines (TC-022).\n"
            "2. Opening an invoice shows its lines grouped by fee type with fee components and GST (TC-019)."
        ),
        "actual": (
            "SCR-02 renders a flat table of individual student_charges rows (Charge description, Gross, Subsidy, Net, Due date). Invoices are only created at checkout, so no INV-YYYY-NNNNNN header, earliest due date aggregation, fee-type grouping, or GST breakdown is shown on SCR-02."
        ),
        "tcs": ["TC-019", "TC-022"],
        "ev_num": 3,
        "caption": "SCR-02 liệt kê từng khoản phí phẳng, không gom theo hóa đơn INV, không chia nhóm Fee Type và không hiển thị GST",
        "ev_file": "evidence_03_BUG-R9-003.png",
        "notes": "Root cause: Architecture/UI mismatch (ASM-R9-UI-01): BillingPaymentPage lists pre-invoice student_charges rather than pre-issued invoices with expandable line groups.",
    },
    {
        "id": "BUG-R9-004",
        "summary": "[SCR-02 Outstanding fees] Charges in a PaymentPending or ReconciliationRequired invoice are hidden instead of shown as Pending and disabled",
        "severity": "S1 (highest)",
        "priority": "Critical",
        "status": "Open",
        "type": "Functional",
        "env": ENV_TEXT,
        "pre": "A student has an open charge that is currently inside a PaymentPending (status = 3) or ReconciliationRequired (status = 5) invoice/transaction.",
        "steps": (
            "1. Start an OnlineGateway checkout for a charge (or set a split payment to ReconciliationRequired) so the invoice is PaymentPending (status = 3).\n"
            "2. Return to SCR-02 (/billing/{studentId}) and inspect the list of items (TC-025, TC-110, TC-127).\n"
            "3. Repeat for a child whose only open invoice is in PaymentPending (TC-034)."
        ),
        "expected": (
            "1. The invoice/item waiting for payment is shown on SCR-02 with status 'Pending' (or 'Verifying') and its selection checkbox is disabled so a second payment cannot be started from the UI (TC-025, TC-110, TC-127).\n"
            "2. A child whose only open invoice is waiting for payment sees that Pending item on SCR-02 and is NOT shown the SCR-03 empty state (TC-034)."
        ),
        "actual": (
            "GetOutstandingChargesHandler filters out any student_charge already linked to a PaymentPending invoice. As a result, the Pending item disappears completely from SCR-02 (TC-025, TC-110, TC-127), and a child whose only open item is PaymentPending is shown the 'No outstanding charges' empty state SCR-03 (TC-034)."
        ),
        "tcs": ["TC-025", "TC-034", "TC-110", "TC-127"],
        "ev_num": 4,
        "caption": "Khoản phí đang PaymentPending / ReconciliationRequired bị ẩn khỏi danh sách SCR-02 và đẩy học sinh chỉ có 1 khoản sang màn rỗng SCR-03",
        "ev_file": "evidence_04_BUG-R9-004.png",
        "notes": "Root cause: GetOutstandingChargesHandler.cs excludes charges joined to non-cancelled invoices instead of returning them with a Pending status flag for BillingPaymentPage to display as disabled.",
    },
    {
        "id": "BUG-R9-005",
        "summary": "[SCR-02 Outstanding fees] Missing sample/demo data notice on the Fees & Payments screen",
        "severity": "S3",
        "priority": "Medium",
        "status": "Open",
        "type": "UI",
        "env": ENV_TEXT,
        "pre": "Parent is signed in and opens SCR-02 (/billing/{studentId}) in the prototype/SIT environment.",
        "steps": (
            "1. Open /billing/01999001-0000-7000-8000-000000000001 (SCR-02).\n"
            "2. Check the page header, banner, and table footer for a notice stating that the fee and balance figures are sample/demo data."
        ),
        "expected": (
            "SCR-02 displays a visible notice labelling the fee and Education Account figures as sample/demo data (TC-027)."
        ),
        "actual": (
            "No sample data, demo data, or prototype notice appears anywhere on SCR-02."
        ),
        "tcs": ["TC-027"],
        "ev_num": 5,
        "caption": "Màn hình SCR-02 không có nhãn thông báo dữ liệu mẫu (Sample / Demo data)",
        "ev_file": "evidence_05_BUG-R9-005.png",
        "notes": "Root cause: BillingPaymentPage.tsx does not include a sample/demo data disclaimer banner.",
    },
    {
        "id": "BUG-R9-006",
        "summary": "[SCR-02 & SCR-04 Student without EA] Missing Education Account (404) shows an error banner and blocks online payment",
        "severity": "S2",
        "priority": "High",
        "status": "Open",
        "type": "Functional",
        "env": ENV_TEXT,
        "pre": "A linked student (e.g., non-citizen or pre-activation student) has outstanding charges in BillingPaymentService but no wallet record in EducationAccountDatabase (wallet API returns 404).",
        "steps": (
            "1. Open SCR-02 (/billing/{studentId}) for a student without an Education Account.\n"
            "2. Observe the balance area and attempt to select a charge and proceed to SCR-04 to pay via Online Gateway."
        ),
        "expected": (
            "1. SCR-02 shows no EA balance and displays a clear note that the student does not have an Education Account (TC-029).\n"
            "2. The Parent can still select outstanding charges and pay via Online Gateway only on SCR-04 (TC-038)."
        ),
        "actual": (
            "When GET /api/v1/accounts/by-student/{studentId} returns HTTP 404, BillingPaymentPage treats it as a fatal load failure, displaying 'Unable to load education account' and hiding the entire charges table and payment flow."
        ),
        "tcs": ["TC-029", "TC-038"],
        "ev_num": 6,
        "caption": "Khi học sinh không có ví EA (API trả 404), UI báo lỗi 'Unable to load education account' và chặn luôn thanh toán Online",
        "ev_file": "evidence_06_BUG-R9-006.png",
        "notes": "Root cause: BillingPaymentPage fails the whole page load when the wallet query returns 404 instead of treating 404 as 'no Education Account' with OnlineGateway-only eligibility.",
    },
    {
        "id": "BUG-R9-007",
        "summary": "[SCR-03 No outstanding fees] Empty state has no in-page link back to the child list",
        "severity": "S2",
        "priority": "High",
        "status": "Open",
        "type": "UI",
        "env": ENV_TEXT,
        "pre": "Parent is signed in and opens /billing/01a15be4-1120-7768-a885-cef71d2a0b8a for Tan Wei Ming (a child with 0 outstanding charges).",
        "steps": (
            "1. Open /billing/01a15be4-1120-7768-a885-cef71d2a0b8a (SCR-03).\n"
            "2. Look inside the main empty-state card ('No outstanding charges') for a button or link to return to the child list (SCR-01)."
        ),
        "expected": (
            "SCR-03 displays the empty state and provides an in-page link/button back to the child list so the Parent can choose another child (TC-032, TC-033)."
        ),
        "actual": (
            "The empty state only renders the text 'No outstanding charges — This student currently has no outstanding education charges to pay.' There is no link or button inside <main> to return to the child list."
        ),
        "tcs": ["TC-032", "TC-033"],
        "ev_num": 7,
        "caption": "Màn hình trống SCR-03 ('No outstanding charges') thiếu nút/liên kết quay về danh sách chọn con trong vùng nội dung chính",
        "ev_file": "evidence_07_BUG-R9-007.png",
        "notes": "Root cause: Empty state block in BillingPaymentPage.tsx omits a 'Back to child list' / 'Choose another child' link.",
    },
    {
        "id": "BUG-R9-008",
        "summary": "[SCR-04 Choose payment method] Closed Education Account is still selectable when wallet balance >= amount",
        "severity": "S2",
        "priority": "High",
        "status": "Open",
        "type": "Functional",
        "env": ENV_TEXT,
        "pre": "A student has an Education Account with status = 'Closed' and balance = SGD 300.00, and an outstanding charge of SGD 75.00.",
        "steps": (
            "1. Open SCR-02 for the student with a Closed EA wallet.\n"
            "2. Select the SGD 75.00 charge and click Continue to open SCR-04 (Choose payment method).\n"
            "3. Inspect the Education Account radio option."
        ),
        "expected": (
            "Because the Education Account is Closed, the Education Account and Split (Both) options are disabled/unavailable, and only Online Gateway can be selected (TC-039)."
        ),
        "actual": (
            "The Education Account radio button is enabled and checked by default because canUseEducationAccount in BillingPaymentPage only compares wallet.balance >= selectedTotal and ignores wallet.status !== 'Active'."
        ),
        "tcs": ["TC-039"],
        "ev_num": 8,
        "caption": "Ví EA có trạng thái Closed nhưng bước chọn phương thức thanh toán (SCR-04) vẫn cho phép chọn Education Account",
        "ev_file": "evidence_08_BUG-R9-008.png",
        "notes": "Root cause: Frontend eligibility helper canUseEducationAccount checks balance >= amount without checking account status == 'Active'.",
    },
    {
        "id": "BUG-R9-009",
        "summary": "[SCR-04 Choose payment method] Low EA balance disables Education Account radio without showing the required explanation message",
        "severity": "S1 (highest)",
        "priority": "Critical",
        "status": "Open",
        "type": "UI",
        "env": ENV_TEXT,
        "pre": "Parent selects charges totalling SGD 470.00 for Test Child Three whose active EA balance is SGD 300.00 (0 < balance < amount to pay).",
        "steps": (
            "1. Open /billing/01999001-0000-7000-8000-000000000003 (Test Child Three, EA balance SGD 300.00).\n"
            "2. Select all charges (SGD 470.00) and click Continue to SCR-04.\n"
            "3. Inspect the disabled Education Account option for the explanation message."
        ),
        "expected": (
            "Education Account is disabled and the screen displays the reason explaining that the Education Account balance is not enough for this amount and suggesting paying the rest online via split payment (TC-041, TC-046)."
        ),
        "actual": (
            "The Education Account radio option is disabled and Both is selected, but no message explains why Education Account is disabled."
        ),
        "tcs": ["TC-041", "TC-046"],
        "ev_num": 9,
        "caption": "SCR-04 khóa nút Education Account khi số dư (300.00) thấp hơn số tiền trả (470.00) nhưng thiếu câu giải thích lý do",
        "ev_file": "evidence_09_BUG-R9-009.png",
        "notes": "Root cause: BillingPaymentPage disables the EducationAccount radio when balance < totalSelectedNet but does not render the low-balance explanation copy from FS 5.1.",
    },
    {
        "id": "BUG-R9-010",
        "summary": "[SCR-04 Split payment] EA portion in split payment is read-only and cannot be lowered or validated on the UI",
        "severity": "S1 (highest)",
        "priority": "Critical",
        "status": "Open",
        "type": "Functional",
        "env": ENV_TEXT,
        "pre": "Parent is on SCR-04 with 'Both (Education Account + Online Gateway)' selected for a charge where 0 < EA balance < total amount.",
        "steps": (
            "1. On SCR-04, select 'Both (Education Account + Online Gateway)'.\n"
            "2. Attempt to edit the Education Account portion (e.g., lower it from 300.00 to 100.00 or 0.01, or test invalid inputs 0.00 and > balance)."
        ),
        "expected": (
            "1. The Parent can lower the EA part (down to 0.01) and the online part updates automatically so both parts add up to the total (TC-051, TC-062).\n"
            "2. Entering 0.00 or an amount above the available balance is refused on SCR-04 with the corresponding validation message (TC-053, TC-057)."
        ),
        "actual": (
            "The Payment split box renders static read-only text ('Use the full available balance from the Education Account first, then pay the remainder online. Education Account SGD 300.00 / Online Gateway SGD 170.00') with no input control."
        ),
        "tcs": ["TC-051", "TC-053", "TC-057", "TC-062"],
        "ev_num": 10,
        "caption": "SCR-04 cố định phần tiền EA trong Split payment ở mức tối đa (read-only), không có ô nhập để phụ huynh chỉnh số tiền EA",
        "ev_file": "evidence_10_BUG-R9-010.png",
        "notes": "Root cause: Frontend implements split payment as read-only full-balance deduction (i18n 'split.readOnlyDescription') even though backend CheckoutPlanner accepts custom eaAmount.",
    },
    {
        "id": "BUG-R9-011",
        "summary": "[API Checkout Split] CheckoutRequestDto ignores onlineAmount and does not reject split requests whose parts sum to less than the invoice total",
        "severity": "S1 (highest)",
        "priority": "Critical",
        "status": "Open",
        "type": "Functional",
        "env": ENV_TEXT,
        "pre": "An outstanding charge of SGD 250.00 exists for a student with an active EA balance of SGD 180.00.",
        "steps": (
            "1. Send POST /api/billing/v1/checkout with paymentMethod = 'Both', eaAmount = 100.00, and onlineAmount = 100.00 (sum = 200.00 < 250.00).\n"
            "2. Inspect the HTTP status code and response payload."
        ),
        "expected": (
            "The API refuses the request with HTTP 400 Bad Request (SPLIT_AMOUNT_INVALID) because the two parts (100.00 + 100.00 = 200.00) do not add up to the 250.00 amount to pay, and no transaction is created (TC-055)."
        ),
        "actual": (
            "The API returns HTTP 200 OK and creates a transaction with EA leg = 100.00 and Online leg = 150.00 because CheckoutRequestDto ignores onlineAmount and silently calculates onlineAmount = total - eaAmount."
        ),
        "tcs": ["TC-055"],
        "ev_num": 11,
        "caption": "API POST /billing/v1/checkout bỏ qua trường onlineAmount nên chấp nhận HTTP 200 khi gửi EA 100 + Online 100 cho khoản phí 250",
        "ev_file": "evidence_11_BUG-R9-011.png",
        "notes": "Root cause: CheckoutRequestDto has no OnlineAmount property; CheckoutPlanner only validates EaAmount and derives the online leg as Total - EaAmount.",
    },
    {
        "id": "BUG-R9-012",
        "summary": "[SCR-05 Review and confirm] Review screen for Education Account payment does not show the remaining EA balance after payment",
        "severity": "S2",
        "priority": "High",
        "status": "Open",
        "type": "UI",
        "env": ENV_TEXT,
        "pre": "Parent selects an outstanding charge on SCR-02, chooses Education Account on SCR-04, and continues to SCR-05 (Check your payment details).",
        "steps": (
            "1. Open /billing/01999001-0000-7000-8000-000000000001.\n"
            "2. Select a charge (e.g., SGD 75.00), keep Education Account selected on Step 2, and click Continue to Step 3 (SCR-05).\n"
            "3. Check the review summary for the remaining Education Account balance after payment."
        ),
        "expected": (
            "SCR-05 shows the child name, invoice/charge details, amount to pay, payment method, and the projected Education Account balance after payment (TC-063)."
        ),
        "actual": (
            "Step 3 ('Check your payment details') shows 'Education Account SGD 75.00', 'Online Gateway SGD 0.00', and the 'Authorize & pay' button, but does not display the remaining EA balance after payment."
        ),
        "tcs": ["TC-063"],
        "ev_num": 12,
        "caption": "Màn hình xem lại trước khi trả tiền (SCR-05) không hiển thị số dư ví Education Account còn lại sau thanh toán",
        "ev_file": "evidence_12_BUG-R9-012.png",
        "notes": "Root cause: Review step in BillingPaymentPage.tsx does not render (wallet.balance - eaAmount) as 'Balance after payment'.",
    },
    {
        "id": "BUG-R9-013",
        "summary": "[Audit Log & Late Webhook] BillingPaymentService has no audit_logs table for F9 payment events and ignores late payment.succeeded webhooks after Failed",
        "severity": "S1 (highest)",
        "priority": "Critical",
        "status": "Open",
        "type": "Data",
        "env": ENV_TEXT,
        "pre": "EA payment, Split payment, or OnlineGateway webhook settlement is executed on BillingPaymentService.",
        "steps": (
            "1. Execute an EA payment or Split payment and query BillingPaymentService for audit log records of F9_PAYMENT_INITIATED, F9_PAYMENT_COMPLETED, and EA reversals (TC-071, TC-076, TC-098, TC-167).\n"
            "2. On a failed OnlineGateway transaction (status = 3 Failed), send a late 'payment.succeeded' webhook with a new eventId and inspect the transaction status."
        ),
        "expected": (
            "1. Every payment initiation, leg transition, completion, and EA reversal is written to the audit log in chronological order.\n"
            "2. A late 'payment.succeeded' webhook arriving after a transaction was already marked Failed transitions the transaction to ReconciliationRequired (status = 5) and writes an alert/audit entry."
        ),
        "actual": (
            "1. BillingPaymentService database has 0 audit log tables (SELECT COUNT(*) FROM sys.tables WHERE name LIKE '%audit%' returns 0) and writes no audit log entries.\n"
            "2. ProcessGatewayWebhookHandler marks a late 'payment.succeeded' webhook on a Failed transaction as Ignored (receipt status = 3) and leaves transaction.status = 3 (Failed) instead of 5 (ReconciliationRequired)."
        ),
        "tcs": ["TC-071", "TC-076", "TC-098", "TC-167"],
        "ev_num": 13,
        "caption": "DB BillingPaymentService thiếu bảng audit_logs cho các sự kiện F9 và bỏ qua (Ignored) webhook payment.succeeded đến muộn sau khi giao dịch đã Failed",
        "ev_file": "evidence_13_BUG-R9-013.png",
        "notes": "Root cause: Audit logging is not implemented in BillingPaymentService, and ProcessGatewayWebhookHandler exits early with WebhookReceiptStatus.Ignored whenever transaction.Status != Pending.",
    },
    {
        "id": "BUG-R9-014",
        "summary": "[Checkout EA & Concurrency] CreateCheckoutHandler lacks EA debit compensation/reversal on save error, timeout reconciliation, and Invoice RowVersion concurrency token",
        "severity": "S1 (highest)",
        "priority": "Critical",
        "status": "Open",
        "type": "Functional",
        "env": ENV_TEXT,
        "pre": "Parent confirms an EA or Split payment in BillingPaymentService.",
        "steps": (
            "1. Inspect CreateCheckoutHandler.cs handling when EducationAccountClient.DebitAsync returns an unclear/timeout error or when saving the invoice/transaction to SQL Server fails after DebitAsync succeeded (TC-073, TC-074, TC-075).\n"
            "2. Inspect InvoiceModelCreating.cs for optimistic concurrency protection on dbo.invoices.row_version (TC-166)."
        ),
        "expected": (
            "1. On unclear EA debit outcome, BillingPaymentService queries Function 1 with the same FundingReference, reverses the debit via CreditAsync if needed, or moves the payment to ReconciliationRequired if reversal fails, showing a safe user message (TC-073, TC-074, TC-075).\n"
            "2. dbo.invoices.row_version is configured as an EF Core concurrency token so concurrent modifications trigger compensation (TC-166)."
        ),
        "actual": (
            "1. CreateCheckoutHandler.cs states 'Happy-path only — no reverse/compensation', propagates raw EA client error messages, and has no compensation block calling CreditAsync if DB persistence fails after wallet debit.\n"
            "2. In InvoiceModelCreating.cs (line 61), '// b.Property(i => i.RowVersion).IsRowVersion();' is commented out."
        ),
        "tcs": ["TC-073", "TC-074", "TC-075", "TC-166"],
        "ev_num": 14,
        "caption": "CreateCheckoutHandler chưa có nhánh hoàn tiền EA (CreditAsync) khi lỗi lưu DB/xung đột đồng thời và InvoiceModelCreating đang tắt IsRowVersion()",
        "ev_file": "evidence_14_BUG-R9-014.png",
        "notes": "Root cause: CreateCheckoutHandler was implemented as happy-path only without compensating CreditAsync calls, and RowVersion concurrency was commented out in InvoiceModelCreating.cs.",
    },
    {
        "id": "BUG-R9-015",
        "summary": "[EF Core Mapping] dbo.invoices student_id and school_id are uniqueidentifier in SQL Server but mapped as string in C#, throwing HTTP 500 InvalidCastException on GetByIdAsync",
        "severity": "S1 (highest)",
        "priority": "Critical",
        "status": "Open",
        "type": "Data",
        "env": ENV_TEXT,
        "pre": "An invoice row exists in BillingPaymentService.dbo.invoices (created by POST /api/billing/v1/checkout).",
        "steps": (
            "1. Create an OnlineGateway or Split (Both) checkout via POST /api/billing/v1/checkout.\n"
            "2. Send a valid signed 'payment.succeeded' webhook to POST /billing/v1/webhooks/payment-gateway (TC-088, TC-089, TC-096, TC-097, TC-115, TC-119, TC-120, TC-143).\n"
            "3. Or resend POST /api/billing/v1/checkout with the exact same Idempotency-Key (up to 100 chars) to replay the first result (TC-078, TC-080)."
        ),
        "expected": (
            "1. On 'payment.succeeded' webhook, SettleSucceededAsync loads the invoice via _invoiceRepo.GetByIdAsync, marks the online leg and transaction Succeeded, marks the invoice Paid (4), and returns HTTP 200 OK.\n"
            "2. On idempotent checkout replay with the same Idempotency-Key, CreateCheckoutHandler loads the existing invoice and returns HTTP 200 OK with the original transaction details."
        ),
        "actual": (
            "Both flows fail with HTTP 500 Internal Server Error: {\"isSuccess\":false,\"statusCode\":500,\"message\":\"Unable to cast object of type 'System.Guid' to type 'System.String'.\"}. Because SettleSucceededAsync throws before committing, no online or split payment can transition to Succeeded/Paid via webhook."
        ),
        "tcs": ["TC-078", "TC-080", "TC-088", "TC-089", "TC-096", "TC-097", "TC-115", "TC-119", "TC-120", "TC-143"],
        "ev_num": 15,
        "caption": "Lỗi nghiêm trọng HTTP 500 InvalidCastException (System.Guid -> System.String) trên dbo.invoices khi xử lý webhook payment.succeeded hoặc gọi lại cùng Idempotency-Key",
        "ev_file": "evidence_15_BUG-R9-015.png",
        "notes": "Root cause: In SQL Server (BillingPaymentService.dbo.invoices), student_id and school_id are uniqueidentifier (Guid), but Invoice.cs and InvoiceModelCreating.cs map StudentId and SchoolId as string (varchar(50)). Fix: add a HasConversion<Guid>() value converter or align the column/property types.",
    },
    {
        "id": "BUG-R9-016",
        "summary": "[API Checkout Validation] Idempotency-Key of 101 characters is not rejected at input validation (returns 404 instead of 400)",
        "severity": "S3",
        "priority": "Medium",
        "status": "Open",
        "type": "Functional",
        "env": ENV_TEXT,
        "pre": "Parent is authenticated and sends a checkout request to POST /api/billing/v1/checkout.",
        "steps": (
            "1. Send POST /api/billing/v1/checkout with idempotencyKey set to a 101-character string ('k' * 101).\n"
            "2. Observe the HTTP response status code."
        ),
        "expected": (
            "The request is rejected immediately at input validation with HTTP 400 Bad Request because Idempotency-Key exceeds the 100-character maximum (TC-081)."
        ),
        "actual": (
            "CheckoutPlanner.ValidateShape does not check IdempotencyKey.Length > 100, so execution proceeds to charge lookup and returns HTTP 404 Not Found ('Charge(s) not found or not outstanding')."
        ),
        "tcs": ["TC-081"],
        "ev_num": 16,
        "caption": "Khóa Idempotency-Key dài 101 ký tự không bị chặn ở bước kiểm tra đầu vào (trả 404 thay vì 400 Bad Request)",
        "ev_file": "evidence_16_BUG-R9-016.png",
        "notes": "Root cause: CheckoutPlanner.ValidateShape only checks string.IsNullOrWhiteSpace(request.IdempotencyKey) and omits request.IdempotencyKey.Trim().Length > 100.",
    },
    {
        "id": "BUG-R9-017",
        "summary": "[Online Payment Failure / Cancel / Expiry] Failed, cancelled or expired online payment leaves invoice stuck in PaymentPending (3) instead of restoring to Issued (2)",
        "severity": "S1 (highest)",
        "priority": "Critical",
        "status": "Open",
        "type": "Functional",
        "env": ENV_TEXT,
        "pre": "An OnlineGateway payment has been initiated (invoice.status = 3 PaymentPending, transaction.status = 1 Pending).",
        "steps": (
            "1. Send a signed 'payment.failed' webhook to POST /billing/v1/webhooks/payment-gateway and inspect dbo.invoices.status and dbo.invoice_items (TC-090, TC-139).\n"
            "2. Send a signed 'payment.cancelled' (TC-091) or 'payment.expired' (TC-092) webhook and inspect dbo.webhook_receipts, dbo.payment_transactions, and dbo.invoices.\n"
            "3. Inspect CreateCheckoutHandler when gateway session creation fails on an OnlineGateway checkout (TC-094)."
        ),
        "expected": (
            "1. On 'payment.failed' or 'payment.expired', the online leg and transaction become Failed (3), the invoice returns to Issued (2), and the charge can be selected again on SCR-02 (TC-090, TC-092, TC-139).\n"
            "2. On 'payment.cancelled', the online leg and transaction become Cancelled (4) and the invoice returns to Issued (2) (TC-091).\n"
            "3. When provider session creation fails during checkout, the invoice is not left in PaymentPending (TC-094)."
        ),
        "actual": (
            "1. On 'payment.failed', SettleFailed sets transaction.status = 3 (Failed) but leaves dbo.invoices.status = 3 (PaymentPending) and keeps dbo.invoice_items attached, permanently locking the charge from SCR-02.\n"
            "2. On 'payment.cancelled' and 'payment.expired', ProcessGatewayWebhookHandler marks the webhook receipt as Ignored (status = 3) and leaves transaction = 1 (Pending) and invoice = 3 (PaymentPending).\n"
            "3. When gateway session creation fails in CreateCheckoutHandler, invoice.Status remains PaymentPending (3)."
        ),
        "tcs": ["TC-090", "TC-091", "TC-092", "TC-094", "TC-139"],
        "ev_num": 17,
        "caption": "Khi thanh toán Online thất bại, bị hủy hoặc hết hạn phiên, hóa đơn bị kẹt ở PaymentPending (status = 3) và không giải phóng khoản phí để thanh toán lại",
        "ev_file": "evidence_17_BUG-R9-017.png",
        "notes": "Root cause: ProcessGatewayWebhookHandler.SettleFailed does not update invoice.Status back to Issued (2) or release invoice_items, and ProcessGatewayWebhookHandler only handles 'payment.succeeded' and 'payment.failed'.",
    },
    {
        "id": "BUG-R9-018",
        "summary": "[SCR-06 Pending & Verifying] Waiting/verifying screen for split payment does not state in prose that the Education Account part was already deducted",
        "severity": "S2",
        "priority": "High",
        "status": "Open",
        "type": "UI",
        "env": ENV_TEXT,
        "pre": "A Split payment (Both: EA SGD 20.00 + Online SGD 30.00) is in Pending status (TC-101) or ReconciliationRequired verifying status (TC-128) on /payment-result (SCR-06).",
        "steps": (
            "1. Open /payment-result?transactionId={txId} for a Pending split payment where the EA leg (SGD 20.00) is Succeeded and the Online leg (SGD 30.00) is Pending (TC-101).\n"
            "2. Open /payment-result?transactionId={txId} for a ReconciliationRequired split payment where the EA leg is also in ReconciliationRequired (TC-128).\n"
            "3. Read the status banner message."
        ),
        "expected": (
            "1. While pending, SCR-06 states in prose that the Education Account portion (SGD 20.00) has already been taken/deducted while the online part is being confirmed (TC-101).\n"
            "2. While verifying (ReconciliationRequired), SCR-06 states the Education Account portion that was taken and offers no retry (TC-128)."
        ),
        "actual": (
            "1. For Pending split, SCR-06 only displays generic text 'Your payment is still being confirmed. This page will update automatically.' without mentioning that the EA part was already deducted.\n"
            "2. For ReconciliationRequired split when EA leg is ReconciliationRequired, SCR-06 only displays 'We are reviewing this payment. Please contact support if the status does not change.' without naming the EA part taken."
        ),
        "tcs": ["TC-101", "TC-128"],
        "ev_num": 18,
        "caption": "Màn hình chờ/đối soát thanh toán gộp (SCR-06) không hiển thị câu thông báo phần tiền Education Account đã bị trừ",
        "ev_file": "evidence_18_BUG-R9-018.png",
        "notes": "Root cause: PaymentResultPage.tsx uses the same generic pending banner for single and split payments, and only shows the EA-deducted note in ReconciliationRequired when eaLeg.status === 'Succeeded'.",
    },
    {
        "id": "BUG-R9-019",
        "summary": "[Split Payment Failure / Cancel / Timeout] Online failure, cancellation or session failure after EA debit does not automatically reverse the EA deduction",
        "severity": "S1 (highest)",
        "priority": "Critical",
        "status": "Open",
        "type": "Functional",
        "env": ENV_TEXT,
        "pre": "A Split payment (Both) has debited the student's Education Account wallet (EA leg = 2 Succeeded) and created an OnlineGateway leg (status = 1 Pending).",
        "steps": (
            "1. Trigger an online failure via 'payment.failed' webhook (TC-103, TC-108, TC-109) or parent cancellation via 'payment.cancelled' webhook (TC-104) on a Split transaction.\n"
            "2. Query BillingPaymentService (payment_transactions, payment_funding_legs) and EducationAccountDatabase (wallets).\n"
            "3. Inspect CreateCheckoutHandler.cs when online session creation fails after EA debit (TC-105) or when online session times out (TC-107)."
        ),
        "expected": (
            "When the online leg of a Split payment fails, is cancelled, fails session creation, or expires, BillingPaymentService calls Function 1 (CreditAsync) to reverse the EA debit, sets the EA leg to Reversed (4) with a ReversalReference, restores the student's wallet balance, marks the transaction Failed/Cancelled, and restores the invoice to Issued (2) (or sets both the EA leg and transaction to ReconciliationRequired (5) if the reversal call itself fails, TC-109)."
        ),
        "actual": (
            "1. ProcessGatewayWebhookHandler.SettleFailed never calls EducationAccountClient.CreditAsync; it sets transaction.Status = 5 (ReconciliationRequired) while leaving the EA leg at status = 2 (Succeeded), reversal_reference = NULL, and wallet.balance = 0.00 (unrefunded).\n"
            "2. In CreateCheckoutHandler.cs (line 226), online session creation failure logs 'TODO reverse EA leg' without calling CreditAsync.\n"
            "3. 'payment.cancelled' is ignored completely."
        ),
        "tcs": ["TC-103", "TC-104", "TC-105", "TC-107", "TC-108", "TC-109"],
        "ev_num": 19,
        "caption": "Khi phần Online của thanh toán gộp (Both) thất bại hoặc bị hủy, hệ thống không gọi hoàn tiền ví EA (số dư ví vẫn bằng 0.00, EA leg giữ Succeeded = 2)",
        "ev_file": "evidence_19_BUG-R9-019.png",
        "notes": "Root cause: SettleFailed in ProcessGatewayWebhookHandler.cs and the fallback branch in CreateCheckoutHandler.cs do not invoke _eaClient.CreditAsync to refund the EA wallet on split payment failure.",
    },
    {
        "id": "BUG-R9-020",
        "summary": "[SCR-07 Receipt] Payment confirmation screen only offers Print receipt and has no Download PDF button or endpoint",
        "severity": "S2",
        "priority": "High",
        "status": "Open",
        "type": "Functional",
        "env": ENV_TEXT,
        "pre": "An Education Account payment (TC-131) or Split payment (TC-132) has succeeded and the Parent is on SCR-07 (/payment-result).",
        "steps": (
            "1. Open /payment-result?transactionId={txId} for a Succeeded EA or Split payment.\n"
            "2. Look for the action to download the receipt as a PDF file."
        ),
        "expected": (
            "SCR-07 provides a 'Download PDF' action that downloads a PDF receipt containing the transaction reference, funding part(s), and total settled amount (TC-131, TC-132)."
        ),
        "actual": (
            "SCR-07 only renders a 'Print receipt' button (which calls window.print()) and a 'Back to home' link. Neither a Download PDF button nor a backend PDF receipt endpoint exists."
        ),
        "tcs": ["TC-131", "TC-132"],
        "ev_num": 20,
        "caption": "Màn hình biên lai thanh toán thành công (SCR-07) chỉ có nút 'Print receipt', thiếu chức năng tải biên lai PDF",
        "ev_file": "evidence_20_BUG-R9-020.png",
        "notes": "Root cause: PaymentResultPage.tsx only implements window.print() and BillingPaymentService has no PDF receipt generation endpoint.",
    },
    {
        "id": "BUG-R9-021",
        "summary": "[SCR-08 Payment failed] Failure screen omits the failure reason, omits the EA-returned notice for refunded splits, and has no Try again button",
        "severity": "S1 (highest)",
        "priority": "Critical",
        "status": "Open",
        "type": "UI",
        "env": ENV_TEXT,
        "pre": "An Online payment or Split payment has failed and the Parent is viewing SCR-08 (/payment-result?transactionId={txId}).",
        "steps": (
            "1. Open /payment-result for a Failed online or split payment that includes a failureReason.\n"
            "2. Check the screen for (a) the short friendly failure reason (TC-136), (b) the statement that the Education Account amount was returned when the EA leg is Reversed (TC-137), and (c) the 'Try again' button to start a new transaction with a new idempotency key (TC-140)."
        ),
        "expected": (
            "1. SCR-08 displays 'Your payment could not be completed. No amount was deducted.', the short failure reason, and the actions 'Try again' and 'Exit' (TC-136, TC-140).\n"
            "2. For a failed split payment whose EA leg was reversed, SCR-08 states 'Your online payment could not be completed. The amount taken from the Education Account has been returned.' (TC-137)."
        ),
        "actual": (
            "SCR-08 only displays 'Your payment could not be completed. Return to Fees & Payments to review the latest outstanding fees.' and a 'Back to home' link. It omits the failure reason, omits the EA-returned message when transaction.status === 'Failed' and eaLeg.status === 'Reversed', and provides no 'Try again' button."
        ),
        "tcs": ["TC-136", "TC-137", "TC-140"],
        "ev_num": 21,
        "caption": "Màn hình thanh toán thất bại (SCR-08) không hiện lý do lỗi, không báo đã hoàn tiền EA cho giao dịch gộp và thiếu nút 'Try again'",
        "ev_file": "evidence_21_BUG-R9-021.png",
        "notes": "Root cause: PaymentResultPage.tsx does not render transaction.failureReason, has no 'Try again' button, and only renders refund/deduction copy under ReconciliationRequired instead of Failed + Reversed.",
    },
    {
        "id": "BUG-R9-022",
        "summary": "[Error Handling UI] 5xx API errors on charge list, checkout confirm, and payment result render raw server error strings (SqlException) on screen",
        "severity": "S2",
        "priority": "High",
        "status": "Open",
        "type": "Security",
        "env": ENV_TEXT,
        "pre": "A 500 Internal Server Error with server error details in errors[] occurs when loading outstanding charges (SCR-02), confirming checkout (SCR-05), or loading payment result (SCR-08).",
        "steps": (
            "1. Trigger or mock an HTTP 500 response containing internal error text (e.g. 'System.Data.SqlClient.SqlException: Timeout expired...') on GET /outstanding-charges (TC-164), POST /checkout (TC-144), or GET /transactions/{id} (TC-138).\n"
            "2. Observe the error message rendered to the Parent on the UI."
        ),
        "expected": (
            "The UI hides all technical details (SQL exceptions, stack traces, class names, HTTP codes) and displays only the safe friendly message 'Something went wrong. Please try again later.' (TC-138, TC-144, TC-164)."
        ),
        "actual": (
            "getBillingPaymentErrorMessage in packages/api-client/src/api-error.ts concatenates and returns response.errors[] verbatim on 5xx responses, causing the UI to display raw 'System.Data.SqlClient.SqlException...' strings directly to the Parent."
        ),
        "tcs": ["TC-138", "TC-144", "TC-164"],
        "ev_num": 22,
        "caption": "Khi API trả lỗi 500 kèm danh sách errors[], giao diện hiển thị nguyên văn chuỗi lỗi kỹ thuật (SqlException) cho phụ huynh",
        "ev_file": "evidence_22_BUG-R9-022.png",
        "notes": "Root cause: packages/api-client/src/api-error.ts returns server errors[] strings even when statusCode >= 500 instead of masking 5xx errors with a generic user-safe message.",
    },
    {
        "id": "BUG-R9-023",
        "summary": "[Database Constraints] dbo.invoice_items allows attaching a student_charge from a different school_id than dbo.invoices.school_id",
        "severity": "S3",
        "priority": "Medium",
        "status": "Open",
        "type": "Data",
        "env": ENV_TEXT,
        "pre": "An invoice exists for School A in BillingPaymentService.dbo.invoices and a student_charge exists for School B in dbo.student_charges.",
        "steps": (
            "1. Insert an invoice with school_id = School A into BillingPaymentService.dbo.invoices.\n"
            "2. Insert a student_charge with school_id = School B into dbo.student_charges.\n"
            "3. Insert a line into dbo.invoice_items linking that student_charge to the invoice (TC-149)."
        ),
        "expected": (
            "The database (and checkout planner) refuses to attach a student_charge whose school_id differs from the invoice's school_id (TC-149)."
        ),
        "actual": (
            "The insert into dbo.invoice_items succeeds ('OK') because neither a composite foreign key / trigger in SQL Server nor a check in CheckoutPlanner verifies that student_charges.school_id matches invoices.school_id."
        ),
        "tcs": ["TC-149"],
        "ev_num": 23,
        "caption": "Bảng dbo.invoice_items cho phép gắn khoản phí (student_charge) của trường khác vào hóa đơn (invoice) mà không từ chối",
        "ev_file": "evidence_23_BUG-R9-023.png",
        "notes": "Root cause: Missing composite FK (invoice_id, school_id) / trigger on dbo.invoice_items and missing multi-school validation in CheckoutPlanner.",
    },
]


def main() -> None:
    tc_to_bug = {}
    for d in DEFECTS:
        for tc in d["tcs"]:
            assert tc not in tc_to_bug, f"Duplicate TC {tc}"
            tc_to_bug[tc] = d["id"]
    print("Total mapped TCs across 23 defects:", len(tc_to_bug))
    assert len(tc_to_bug) == 65, f"Expected 65 TCs, got {len(tc_to_bug)}"

    src_path = r"D:\AIQA\Agileproject\roles\qa\R9_Test_Cases_final_FS09_153cases.xlsx"
    out_def_path = r"D:\AIQA\Agileproject\roles\qa\R9_Test_Cases_final_FS09_defects.xlsx"
    orig_path = r"D:\AIQA\Agileproject\roles\qa\R9_Test_Cases_final_FS09.xlsx"

    wb = openpyxl.load_workbook(src_path)
    ws_tc = wb["Test cases "]

    font_bug = Font(name="Calibri", size=10, bold=True, color="B91C1C")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    mapped_in_sheet = 0
    for r in range(2, ws_tc.max_row + 1):
        tc_id = ws_tc.cell(r, 1).value
        if not tc_id or not str(tc_id).strip().startswith("TC-"):
            continue
        tc_id = str(tc_id).strip()
        bug_cell = ws_tc.cell(r, 17)
        if tc_id in tc_to_bug:
            bug_cell.value = tc_to_bug[tc_id]
            bug_cell.font = font_bug
            bug_cell.alignment = align_center
            mapped_in_sheet += 1
        else:
            bug_cell.value = None

    print("Populated Defect ID in Test cases sheet for:", mapped_in_sheet, "rows")

    ws_dr = wb["Defect Reports"]
    thin_border = Border(
        left=Side(style="thin", color="D0D7DE"),
        right=Side(style="thin", color="D0D7DE"),
        top=Side(style="thin", color="D0D7DE"),
        bottom=Side(style="thin", color="D0D7DE"),
    )
    font_regular = Font(name="Calibri", size=10, bold=False, color="1F2328")
    font_id = Font(name="Calibri", size=10, bold=True, color="1F2328")
    align_left = Alignment(horizontal="left", vertical="top", wrap_text=True)
    align_mid = Alignment(horizontal="center", vertical="top", wrap_text=True)

    sev_fills = {
        "S1 (highest)": PatternFill("solid", fgColor="FEE2E2"),
        "S2": PatternFill("solid", fgColor="FFEDD5"),
        "S3": PatternFill("solid", fgColor="FEF9C3"),
        "S4 (lowest)": PatternFill("solid", fgColor="F3F4F6"),
    }
    status_fill = PatternFill("solid", fgColor="DBEAFE")

    for r in range(2, max(ws_dr.max_row + 1, 26)):
        for c in range(1, 15):
            ws_dr.cell(r, c).value = None

    for idx, d in enumerate(DEFECTS, start=2):
        row_vals = [
            d["id"],
            d["summary"],
            d["severity"],
            d["priority"],
            d["status"],
            d["type"],
            d["env"],
            d["pre"],
            d["steps"],
            d["expected"],
            d["actual"],
            ", ".join(d["tcs"]),
            f"Image {d['ev_num']} (sheet Evidence)\n{d['ev_file']}",
            d["notes"],
        ]
        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws_dr.cell(idx, c_idx, val)
            cell.border = thin_border
            cell.font = font_id if c_idx == 1 else font_regular
            cell.alignment = align_mid if c_idx in (1, 3, 4, 5, 6) else align_left
            if c_idx == 3 and val in sev_fills:
                cell.fill = sev_fills[val]
            elif c_idx == 5:
                cell.fill = status_fill
        ws_dr.row_dimensions[idx].height = 95

    col_widths = {
        "A": 16, "B": 36, "C": 14, "D": 12, "E": 12, "F": 16,
        "G": 26, "H": 28, "I": 42, "J": 38, "K": 42, "L": 24, "M": 26, "N": 38,
    }
    for col_l, w in col_widths.items():
        ws_dr.column_dimensions[col_l].width = w

    ws_dr.data_validations.dataValidation.clear()
    dv_sev = DataValidation(type="list", formula1='"S1 (highest),S2,S3,S4 (lowest)"', allow_blank=True)
    dv_pri = DataValidation(type="list", formula1='"Critical,High,Medium,Low"', allow_blank=True)
    dv_sta = DataValidation(type="list", formula1='"Open,In Progress,Resolved,Closed,Rejected,Deferred"', allow_blank=True)
    ws_dr.add_data_validation(dv_sev)
    ws_dr.add_data_validation(dv_pri)
    ws_dr.add_data_validation(dv_sta)
    dv_sev.add("C2:C30")
    dv_pri.add("D2:D30")
    dv_sta.add("E2:E30")

    if "Evidence" in wb.sheetnames:
        del wb["Evidence"]
    dr_index = wb.sheetnames.index("Defect Reports")
    ws_ev = wb.create_sheet(title="Evidence", index=dr_index + 1)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    headers_ev = ["Image", "Caption from tester", "Defect ID", "Screenshot"]
    for c_idx, h_text in enumerate(headers_ev, start=1):
        cell = ws_ev.cell(1, c_idx, h_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws_ev.row_dimensions[1].height = 26
    ws_ev.freeze_panes = "A2"

    ws_ev.column_dimensions["A"].width = 10
    ws_ev.column_dimensions["B"].width = 48
    ws_ev.column_dimensions["C"].width = 22
    ws_ev.column_dimensions["D"].width = 98

    for idx, d in enumerate(DEFECTS, start=2):
        c1 = ws_ev.cell(idx, 1, d["ev_num"])
        c2 = ws_ev.cell(idx, 2, f"{d['caption']}\n(Linked: {', '.join(d['tcs'])})")
        c3 = ws_ev.cell(idx, 3, d["id"])
        c4 = ws_ev.cell(idx, 4, None)
        c1.alignment = align_mid
        c2.alignment = align_left
        c3.alignment = align_mid
        c1.font = font_id
        c2.font = font_regular
        c3.font = font_bug
        for cell in (c1, c2, c3, c4):
            cell.border = thin_border

        img_file = EV_DIR / d["ev_file"]
        if img_file.exists():
            xl_img = XLImage(str(img_file))
            scale = 660.0 / float(xl_img.width)
            xl_img.width = int(xl_img.width * scale)
            xl_img.height = int(xl_img.height * scale)
            ws_ev.add_image(xl_img, f"D{idx}")
            ws_ev.row_dimensions[idx].height = max(180, int(xl_img.height * 0.78) + 14)

    ws_sum = wb["Summary"]
    ws_sum["A66"] = (
        "Execution Status and Actual Result reflect the full 153/153 automated run of 22 Oct 2026 "
        "(88 Passed, 65 Failed, 0 Not Run). All 65 Failed cases are linked in Column Q (Defect ID) "
        "to BUG-R9-001 through BUG-R9-023 in the Defect Reports and Evidence sheets."
    )

    wb.save(out_def_path)
    print("Saved complete workbook with Defects & Evidence to:", out_def_path)
    for target in (src_path, orig_path):
        try:
            wb.save(target)
            print("Also saved directly to:", target)
        except PermissionError:
            print("Target file is currently open in Excel/WPS (locked):", target)


if __name__ == "__main__":
    main()
