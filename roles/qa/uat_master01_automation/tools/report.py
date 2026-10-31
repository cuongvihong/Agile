"""Build the reports of the last run from artifacts/results.json.

  artifacts/uat_summary.csv          one line per case
  artifacts/uat_steps.csv            one line per step
  artifacts/UAT_Master01_Report.xlsx sheets Summary, Defects (candidate defects with evidence), Steps, Data mapping

A defect candidate is a step that is FAIL because a check is false (wrong behaviour) or GAP (the build lacks what the
master describes; approved as FAIL on 30/10/2026). A FAIL with kind ERROR means the step could not be driven: look at
the screenshot and video first, it can be a product defect or a script problem.
"""
from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from core import config  # noqa: E402

ORDER = ["UAT-MOE-01", "UAT-MOE-04", "UAT-MOE-05", "UAT-MOE-05-E1", "UAT-MOE-06", "UAT-MOE-06-N1",
         "UAT-SCH-01", "UAT-SCH-01-N1", "UAT-SCH-01-E1", "UAT-SCH-02", "UAT-SCH-02-N1", "UAT-SCH-03",
         "UAT-SCH-03-N1", "UAT-SCH-04", "UAT-SCH-04-E1", "UAT-PAR-01", "UAT-PAR-01-N1", "UAT-PAR-01-N2",
         "UAT-PAR-01-E1", "UAT-PAR-01-N3", "UAT-PAR-01-A1", "UAT-PAR-01-N4", "UAT-PAR-01-E2", "UAT-PAR-02",
         "UAT-PAR-02-N1", "UAT-PAR-02-N2", "UAT-PAR-02-E1", "UAT-PAR-02-E2", "UAT-PAR-02-A1", "UAT-PAR-02-N3",
         "UAT-PAR-02-A2", "UAT-PAR-02-N4", "UAT-PAR-03", "UAT-PAR-03-A1", "UAT-PAR-03-E1"]


def load() -> dict:
    try:
        return json.loads(config.RESULTS_JSON.read_text(encoding="utf-8"))
    except Exception:
        return {}


def reason(step: dict) -> str:
    bad = [f"NOT OK {c['label']} (actual: {c['actual']})" for c in step["checks"] if not c["ok"]]
    parts = bad + [f"GAP {g}" for g in step["gaps"]]
    if step["error"]:
        parts.append(f"ERROR {step['error']}")
    if step["blocked"]:
        parts.append(f"BLOCKED {step['blocked']}")
    if step.get("api_errors"):
        parts.append("API " + "; ".join(step["api_errors"][:3]))
    return " | ".join(parts)


def main() -> None:
    res = load()
    cases = [res[c] for c in ORDER if c in res] + [v for k, v in res.items() if k not in ORDER]
    summary_rows, step_rows, defects, mapping = [], [], [], []
    for c in cases:
        bad_steps = [s for s in c["steps"] if s["status"] != "PASS"]
        equiv = [e for s in c["steps"] for e in s.get("equiv", [])]
        summary_rows.append({
            "Case": c["case"], "Title": c["title"], "Result": c["result"], "Severity 1": "Yes" if c.get("severity1") else "",
            "Steps": len(c["steps"]), "Not passed": "; ".join(f"step {s['n']} {s['status']}" for s in bad_steps),
            "Naming differences (EQUIV)": " | ".join(equiv), "Ran at": c["ran_at"], "Videos": " ".join(c.get("videos", [])),
        })
        for m in c.get("data_mapping", []):
            mapping.append({"Case": c["case"], "Mapping": m})
        for s in c["steps"]:
            step_rows.append({
                "Case": c["case"], "Step": s["n"], "Title": s["title"], "Expected": s["expected"], "Status": s["status"],
                "Details": reason(s) or "; ".join(f"OK {k['label']}" for k in s["checks"]),
                "Notes": " | ".join(s.get("notes", [])), "Screenshots": " ".join(s.get("shots", [])),
            })
            wrong = [k for k in s["checks"] if not k["ok"]]
            if s["status"] in ("FAIL", "GAP"):
                kind = "WRONG BEHAVIOUR" if wrong else ("GAP (not in build)" if s["gaps"] else "ERROR (check evidence)")
                defects.append({
                    "No": len(defects) + 1, "Case": c["case"], "Step": s["n"], "Kind": kind,
                    "Severity": "High" if c.get("severity1") else "Medium",
                    "Summary": (wrong[0]["label"] if wrong else (s["gaps"][0] if s["gaps"] else s["error"]))[:180],
                    "Step title": s["title"], "Expected (master 01)": s["expected"],
                    "Actual": reason(s), "Screenshot": (s.get("shots") or [""])[0], "Video": " ".join(c.get("videos", [])),
                })
    _csv(config.ARTIFACTS / "uat_summary.csv", summary_rows)
    _csv(config.ARTIFACTS / "uat_steps.csv", step_rows)
    _xlsx(summary_rows, defects, step_rows, mapping)
    counts = {}
    for r in summary_rows:
        counts[r["Result"]] = counts.get(r["Result"], 0) + 1
    print(f"{len(summary_rows)} case(s): {counts}. Defect candidates: {len(defects)}")
    print(f"Report: {config.ARTIFACTS / 'UAT_Master01_Report.xlsx'}")


def _csv(path: Path, rows: list[dict]) -> None:
    if not rows:
        return
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)


def _xlsx(summary, defects, steps, mapping) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    fills = {"PASS": "C6EFCE", "FAIL": "FFC7CE", "GAP": "E4D0F5", "BLOCKED": "FFE5B4"}
    wb = Workbook()
    for title, rows in (("Summary", summary), ("Defects", defects), ("Steps", steps), ("Data mapping", mapping)):
        ws = wb.active if title == "Summary" else wb.create_sheet(title)
        ws.title = title
        if not rows:
            ws.append(["(none)"])
            continue
        heads = list(rows[0])
        ws.append(heads)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3A93")
        for r in rows:
            ws.append([r[h] for h in heads])
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if heads[cell.column - 1] in ("Result", "Status") and cell.value in fills:
                    cell.fill = PatternFill("solid", fgColor=fills[cell.value])
        for i, h in enumerate(heads, start=1):
            width = 60 if h in ("Actual", "Details", "Expected (master 01)", "Expected", "Summary", "Mapping", "Notes", "Not passed") else 18
            ws.column_dimensions[ws.cell(1, i).column_letter].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(config.ARTIFACTS / "UAT_Master01_Report.xlsx")


if __name__ == "__main__":
    main()
